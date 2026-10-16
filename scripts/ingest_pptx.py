#!/usr/bin/env python3
"""已有 PPTX 解析（美化路由第 1 步 · Ingest）→ source.json

逐页抽取：标题、文本段（含层级、字号、颜色、位置）、表格、图表（类型 + 类别 + 系列数据）、图片、备注；
SmartArt / OLE / 媒体等不支持的对象记录到 unsupported[]（不重绘、不猜内容），由美化报告逐项列出。

用法：python scripts/ingest_pptx.py in.pptx [--out source.json] [--media-dir DIR]
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.util import Emu

EMU_IN = 914400
CHART_KIND = {  # python-pptx XL_CHART_TYPE 名 → 我们的 chart.type
    'COLUMN': 'bar', 'BAR': 'hbar', 'LINE': 'line', 'AREA': 'area', 'PIE': 'donut', 'DOUGHNUT': 'donut',
    'XY_SCATTER': 'scatter', 'BUBBLE': 'bubble', 'RADAR': 'radar',
}


def _in(v) -> float:
    return round((v or 0) / EMU_IN, 3)


def _chart_type(ch) -> str:
    name = str(ch.chart_type).split('.')[-1].split(' ')[0]
    for k, v in CHART_KIND.items():
        if name.startswith(k):
            if v == 'bar' and 'STACKED' in name:
                return 'stack'
            return v
    return 'bar'


def _para_info(p):
    text = ''.join(r.text for r in p.runs) if p.runs else p.text
    sizes = sorted({r.font.size.pt for r in p.runs if r.font.size is not None})
    colors = []
    for r in p.runs:
        try:
            if r.font.color and r.font.color.type is not None and r.font.color.rgb is not None:
                colors.append(str(r.font.color.rgb))
        except (AttributeError, TypeError):
            pass
    return {'text': text, 'level': p.level or 0, 'sizes': sizes, 'colors': sorted(set(colors)),
            'bold': any(bool(r.font.bold) for r in p.runs)}


def _walk(shapes, out, media_dir, slide_no, unsupported):
    for sh in shapes:
        box = {'x': _in(sh.left), 'y': _in(sh.top), 'w': _in(sh.width), 'h': _in(sh.height)}
        st = sh.shape_type
        if st == MSO_SHAPE_TYPE.GROUP:
            _walk(sh.shapes, out, media_dir, slide_no, unsupported)
            continue
        if getattr(sh, 'has_chart', False) and sh.has_chart:
            ch = sh.chart
            cats = [str(c) for c in ch.plots[0].categories] if ch.plots else []
            series = []
            for pl in ch.plots:
                for s in pl.series:
                    series.append({'name': s.name, 'values': [None if v is None else float(v) for v in s.values]})
            out.append({'kind': 'chart', 'chartType': _chart_type(ch), 'categories': cats, 'series': series,
                        'title': ch.chart_title.text_frame.text if ch.has_title and ch.chart_title.has_text_frame else '',
                        **box})
            continue
        if getattr(sh, 'has_table', False) and sh.has_table:
            rows = [[c.text for c in r.cells] for r in sh.table.rows]
            out.append({'kind': 'table', 'rows': rows, **box})
            continue
        if st == MSO_SHAPE_TYPE.PICTURE:
            img = sh.image
            name = f's{slide_no:02d}-{hashlib.sha1(img.blob).hexdigest()[:10]}.{img.ext}'
            if media_dir:
                os.makedirs(media_dir, exist_ok=True)
                with open(os.path.join(media_dir, name), 'wb') as fh:
                    fh.write(img.blob)
            out.append({'kind': 'picture', 'file': name, 'alt': sh.name, **box})
            continue
        if st in (MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT, MSO_SHAPE_TYPE.LINKED_OLE_OBJECT, MSO_SHAPE_TYPE.MEDIA) or \
                sh.element.tag.endswith('graphicFrame') and not getattr(sh, 'has_text_frame', False):
            unsupported.append({'slide': slide_no, 'shape': sh.name, 'type': str(st), 'action': '无法解析，未迁入新版式；原稿该页保留原对象，需人工处理'})
            continue
        if getattr(sh, 'has_text_frame', False) and sh.has_text_frame:
            paras = [_para_info(p) for p in sh.text_frame.paragraphs]
            paras = [p for p in paras if p['text'].strip()]
            if not paras:
                continue
            is_title = False
            if sh.is_placeholder:
                try:
                    is_title = sh.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
                except ValueError:
                    pass
            out.append({'kind': 'text', 'isTitle': is_title, 'paras': paras, **box})


def ingest(path: str, media_dir: str | None = None) -> dict:
    prs = Presentation(path)
    slides, unsupported = [], []
    for i, sl in enumerate(prs.slides, 1):
        shapes: list = []
        _walk(sl.shapes, shapes, media_dir, i, unsupported)
        shapes.sort(key=lambda s: (round(s['y'], 1), s['x']))
        texts = [s for s in shapes if s['kind'] == 'text']
        # 标题：占位符标题优先；否则取页面上方 1/4 内最大字号的文本框
        tshape = next((s for s in texts if s['isTitle']), None)
        if tshape is None and texts:
            top = [s for s in texts if s['y'] < 1.6 and len(s['paras']) == 1]
            if top:
                tshape = max(top, key=lambda s: max(s['paras'][0]['sizes'] or [0]))
        notes = ''
        if sl.has_notes_slide:
            notes = sl.notes_slide.notes_text_frame.text if sl.notes_slide.notes_text_frame else ''
        slides.append({
            'no': i,
            'title': tshape['paras'][0]['text'] if tshape else '',
            'titleShape': shapes.index(tshape) if tshape else None,
            'shapes': shapes,
            'notes': notes,
        })
    return {'source': os.path.basename(path), 'slideW': _in(prs.slide_width), 'slideH': _in(prs.slide_height),
            'slides': slides, 'unsupported': unsupported}


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument('pptx')
    ap.add_argument('--out')
    ap.add_argument('--media-dir')
    a = ap.parse_args()
    data = ingest(a.pptx, a.media_dir)
    js = json.dumps(data, ensure_ascii=False, indent=1)
    if a.out:
        with open(a.out, 'w', encoding='utf-8') as fh:
            fh.write(js)
    else:
        sys.stdout.write(js)
    return 0


if __name__ == '__main__':
    sys.exit(main())
